"""Print blocks (see cvprint) to a Word file (.docx), in memory.

Built the way applicant-tracking systems read best: one column, Arial, Word's own heading style
for section headings and its own list style for bullets, dates set against a right-hand tab stop.
No tables, text boxes, columns or images, and nothing in the page header or footer.

Nothing is written to disk."""
import io
from datetime import datetime, timezone

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

import cvprint as page

FONT = "Arial"
HEADING_STYLE, BULLET_STYLE = "Heading 1", "List Bullet"
BLACK = RGBColor(0, 0, 0)
# what may follow a paragraph border inside a paragraph's properties, in the order Word's file format demands
AFTER_BORDER = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
                "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
                "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def _restyle(style, size: float, *, bold: bool = False, before: float = 0, after: float = 0, keep_with_next: bool = False) -> None:
    """Make one of Word's built-in styles plain: Arial, black, single spaced. The template's styles
    take their typeface from the document theme, and that link has to be cut for Arial to hold."""
    style.font.name, style.font.size, style.font.bold, style.font.italic, style.font.color.rgb = FONT, Pt(size), bold, False, BLACK
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for script in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{script}"), FONT)
    for theme in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{theme}"), None)
    spacing = style.paragraph_format
    spacing.space_before, spacing.space_after, spacing.line_spacing, spacing.keep_with_next = Pt(before), Pt(after), 1.0, keep_with_next


def _rule(paragraph) -> None:
    """A thin line under a section heading: a paragraph border, not a table or a drawing."""
    border, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
    for name, value in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "000000")):
        bottom.set(qn(f"w:{name}"), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().insert_element_before(border, *AFTER_BORDER)


def render(blocks: list[tuple], *, title: str, author: str) -> bytes:
    """The blocks as an A4 Word file."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(page.PAGE_MM[0]), Mm(page.PAGE_MM[1])
    section.left_margin = section.right_margin = Mm(page.MARGIN_X_MM)
    section.top_margin = section.bottom_margin = Mm(page.MARGIN_Y_MM)
    text_width = section.page_width - section.left_margin - section.right_margin
    _restyle(doc.styles["Normal"], page.BODY_PT)
    _restyle(doc.styles[HEADING_STYLE], page.HEADING_PT, bold=True, before=page.HEADING_BEFORE_PT, after=page.HEADING_AFTER_PT, keep_with_next=True)
    _restyle(doc.styles[BULLET_STYLE], page.BODY_PT)
    bullets = doc.styles[BULLET_STYLE].paragraph_format
    bullets.left_indent, bullets.first_line_indent = Pt(page.BULLET_INDENT_PT), Pt(-page.BULLET_INDENT_PT)
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    properties = doc.core_properties  # the template's own say "python-docx" and a date in 2013
    properties.author = properties.last_modified_by = page.clean(author)
    properties.title, properties.comments, properties.created, properties.modified = page.clean(title), "", now, now

    for block in blocks:
        kind, text = block[0], page.clean(block[1])
        if kind == "gap":
            doc.add_paragraph()
        elif not text:
            continue
        elif kind == "name":
            run = doc.add_paragraph().add_run(text)
            run.bold, run.font.size = True, Pt(page.NAME_PT)
        elif kind == "heading":
            _rule(doc.add_paragraph(text, style=HEADING_STYLE))
        elif kind == "entry":
            paragraph = doc.add_paragraph()
            layout = paragraph.paragraph_format
            layout.space_before, layout.keep_with_next = Pt(page.ENTRY_BEFORE_PT), True
            paragraph.add_run(text).bold = True
            dates = page.clean(block[2]) if len(block) > 2 else ""
            if dates:  # dates sit at the right margin by a tab stop, never a table
                layout.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
                paragraph.add_run(f"\t{dates}")
        elif kind == "sub":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run(text).bold = True
        elif kind == "bullet":
            doc.add_paragraph(text, style=BULLET_STYLE)
        elif kind == "para":
            doc.add_paragraph(text).paragraph_format.space_after = Pt(page.PARA_AFTER_PT)
        else:  # contact, line
            doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
