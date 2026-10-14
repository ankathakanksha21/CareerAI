"""The CV as a Word file, read back with python-docx. No database; nothing may be written to disk."""
import io
import zipfile
from datetime import date

import pytest
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn

import cvpdf
import cvprint
import cvword
from cvsample import BLOCKS, content, missing_or_out_of_order, pdf_text, pieces, word, word_text
from files import forbid_disk_writes

WHO = {"title": "Zarvin Quillfeather CV", "author": "Zarvin Quillfeather"}


def render(blocks):
    return cvword.render(blocks, **WHO)


def part(data: bytes, name: str) -> str:
    return zipfile.ZipFile(io.BytesIO(data)).read(name).decode("utf-8")


def test_every_block_is_there_in_order():
    data = render(BLOCKS)
    assert data[:2] == b"PK"
    assert missing_or_out_of_order(pieces(BLOCKS), word_text(data)) == []
    assert len([p for p in word(data).paragraphs if p.text.strip()]) == len(BLOCKS)  # one paragraph to a block, nothing more


def test_no_tables_and_nothing_in_header_or_footer():
    data = render(BLOCKS)
    doc = word(data)
    assert doc.tables == [] and len(doc.sections) == 1 and doc.inline_shapes._inline_lst == []
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert not [n for n in names if n.startswith(("word/header", "word/footer", "word/media", "word/embeddings"))]
    body = part(data, "word/document.xml")
    for banned in ("<w:tbl", "<w:drawing", "<w:pict", "w:txbxContent", "<w:headerReference", "<w:footerReference", "<w:cols w:num"):
        assert banned not in body, banned


def test_headings_use_the_heading_style_and_bullets_the_list_style():
    doc = word(render(BLOCKS))
    by_style = {}
    for p in doc.paragraphs:
        by_style.setdefault(p.style.name, []).append(p.text)
    assert by_style["Heading 1"] == [b[1] for b in BLOCKS if b[0] == "heading"]
    assert by_style["List Bullet"] == [b[1] for b in BLOCKS if b[0] == "bullet"]
    assert set(by_style) == {"Normal", "Heading 1", "List Bullet"}
    bullet = next(p for p in doc.paragraphs if p.style.name == "List Bullet")
    assert bullet.style.element.pPr.numPr is not None  # a real list, not a typed bullet character
    assert not any(p.text.lstrip().startswith(("•", "-", "*")) for p in doc.paragraphs)


def test_dates_follow_a_tab_not_a_table():
    doc = word(render(BLOCKS))
    entry = next(p for p in doc.paragraphs if p.text.startswith("Zzqx Holdings Ltd"))
    assert entry.text == "Zzqx Holdings Ltd, Pune, India\tJul 2023 to Aug 2025"
    stops = list(entry.paragraph_format.tab_stops)
    assert len(stops) == 1 and stops[0].alignment == WD_TAB_ALIGNMENT.RIGHT
    section = doc.sections[0]
    assert stops[0].position == section.page_width - section.left_margin - section.right_margin
    assert entry.runs[0].bold is True and entry.runs[0].text == "Zzqx Holdings Ltd, Pune, India"
    undated = next(p for p in doc.paragraphs if p.text.startswith("Zzqx Foods"))
    assert undated.text == "Zzqx Foods" and list(undated.paragraph_format.tab_stops) == []


def test_page_is_a4_and_font_is_arial():
    doc = word(render(BLOCKS))
    section = doc.sections[0]
    assert (round(section.page_width.mm), round(section.page_height.mm)) == (210, 297)
    assert round(section.left_margin.mm) == round(section.right_margin.mm) == 18 and round(section.top_margin.mm) == 16
    for name, size in (("Normal", 10.5), ("Heading 1", 11.5), ("List Bullet", 10.5)):
        style = doc.styles[name]
        fonts = style.element.rPr.rFonts
        assert style.font.name == "Arial" and style.font.size.pt == size, name
        assert {fonts.get(qn(f"w:{k}")) for k in ("ascii", "hAnsi", "eastAsia", "cs")} == {"Arial"}, name
        assert not [k for k in fonts.attrib if k.endswith("Theme") or k.endswith("cstheme")], name  # no link back to the theme's typeface
        assert str(style.font.color.rgb) == "000000" and style.paragraph_format.line_spacing == 1.0, name
    name_run = doc.paragraphs[0].runs[0]
    assert name_run.bold is True and name_run.font.size.pt == 16


@pytest.mark.parametrize("text", ["Nguyễn Thị Çağla Łukasz-O'Brien", "Московский государственный университет", "Ελληνικά (Basic)",
                                  "Saved £120,000 and ₹5 lakh", "中文 and العربية and हिन्दी and 😀"])
def test_accents_other_scripts_and_a_pound_sign_survive(text):
    assert text in word_text(render([("name", text), ("line", text), ("bullet", text)]))


def test_ampersand_angle_brackets_and_a_control_character():
    data = render([("line", "R&D <team> \x07lead\x00 at \"Zzqx\" & 'Sons'"), ("entry", "A & B\x0b Ltd", "Jul <2023>"), ("bullet", "x\ud800y￾")])
    assert [p.text for p in word(data).paragraphs] == ["R&D <team> lead at \"Zzqx\" & 'Sons'", "A & B Ltd\tJul <2023>", "xy"]


def test_author_is_the_user_not_the_library():
    data = render(BLOCKS)
    properties = word(data).core_properties
    assert (properties.author, properties.last_modified_by, properties.title) == ("Zarvin Quillfeather", "Zarvin Quillfeather", "Zarvin Quillfeather CV")
    assert properties.created.year >= 2026 and properties.comments == ""
    assert "python-docx" not in part(data, "docProps/core.xml")


def test_empty_blocks_and_blank_text_are_skipped():
    doc = word(render([("line", ""), ("bullet", "  "), ("entry", "", "Jul 2023"), ("gap", ""), ("line", "Only this")]))
    assert [p.text for p in doc.paragraphs] == ["", "Only this"]


def test_no_file_is_written(monkeypatch):
    forbid_disk_writes(monkeypatch)
    assert render(BLOCKS)[:2] == b"PK"


def test_word_and_pdf_say_the_same_thing_in_the_same_order():
    c = content(right_to_work={"shown": True, "text": "Visa status: Student visa, valid until 30 January 2027."})
    c["work"][0]["bullets"][1]["hidden"] = True
    for blocks in (cvprint.blocks(c), cvprint.letter_blocks(c, "Ms Patel", date(2026, 10, 13))):
        wanted = pieces(blocks)
        in_word, in_pdf = word_text(cvword.render(blocks, **WHO)), pdf_text(cvpdf.render(blocks, **WHO)[0])
        assert missing_or_out_of_order(wanted, in_word) == [] and missing_or_out_of_order(wanted, in_pdf) == []
        assert "Automated a three-spreadsheet merge" not in in_word + in_pdf  # the hidden bullet is in neither file
        assert in_pdf.replace("• ", "") == in_word  # and neither file says anything the other does not
